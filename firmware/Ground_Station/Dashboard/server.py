import serial
import serial.tools.list_ports
import threading
import json
import asyncio
import websockets
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import webbrowser
import os
import time
import math

# ── CONFIG ──────────────────────────────────────────────────────────────────
PORT    = "COM7"    # Change to your port, e.g. "/dev/ttyUSB0" on Linux/Mac
BAUD    = 115200
WS_PORT = 8765

# ── STATE ────────────────────────────────────────────────────────────────────
data_log   = []       # list of telemetry dicts
sos_events = []       # list of SOS_EVENT dicts
latest_sos = None     # most recent SOS_EVENT dict (sent to new clients on connect)
connected_clients = set()
lock = threading.Lock()
loop = None

def get_timestamp():
    return datetime.now().strftime("%H:%M:%S")

# ── WEBSOCKET BROADCAST ──────────────────────────────────────────────────────
async def broadcast(msg):
    with lock:
        clients = set(connected_clients)
    if clients:
        data = json.dumps(msg)
        await asyncio.gather(*[c.send(data) for c in clients], return_exceptions=True)

def broadcast_sync(msg):
    if loop and loop.is_running():
        asyncio.run_coroutine_threadsafe(broadcast(msg), loop)

# ── SERIAL PARSER ────────────────────────────────────────────────────────────
#
#  Ground station outputs exactly two line types:
#
#  1. CANSAT,temp,hum,pres,alt,lat,lon
#       [0]    [1]  [2] [3]  [4] [5] [6]
#
#  2. SOS_EVENT,ctemp,chum,cpres,calt,clat,clon,bpm,slat,slon,stemp,shum
#       [0]       [1]   [2]  [3]  [4]  [5]  [6] [7]  [8]  [9]  [10]  [11]
#
#  Everything else (boot messages, [ERR], ---) is ignored.
# ────────────────────────────────────────────────────────────────────────────
def parse_and_dispatch(line, now):
    global latest_sos

    line = line.strip()
    if not line:
        return

    # ── CANSAT telemetry ─────────────────────────────────────────────────────
    if line.startswith("CANSAT,"):
        parts = line.split(",")
        if len(parts) < 7:
            print(f"[WARN] Short CANSAT line: {line}")
            return
        try:
            temp     = float(parts[1])
            humidity = float(parts[2])
            pressure = float(parts[3])
            altitude = float(parts[4])
            lat      = float(parts[5]) if parts[5].upper() != "NONE" else None
            lon      = float(parts[6]) if parts[6].upper() != "NONE" else None
        except ValueError:
            print(f"[WARN] Could not parse CANSAT line: {line}")
            return

        entry = {
            "time": now, "temp": round(temp, 2), "humidity": round(humidity, 2),
            "pressure": round(pressure, 2), "altitude": round(altitude, 2),
            "lat": lat, "lon": lon,
        }
        with lock:
            data_log.append(entry)
        broadcast_sync({"type": "telemetry", **entry})
        return

    # ── SOS_EVENT ────────────────────────────────────────────────────────────
    if line.startswith("SOS_EVENT,"):
        parts = line.split(",")
        if len(parts) < 12:
            print(f"[WARN] Short SOS_EVENT line: {line}")
            return

        def safe_float(v):
            try:
                return float(v) if v.upper() != "NONE" else None
            except ValueError:
                return None

        def safe_str(v):
            return None if v.upper() == "NONE" else v.strip()

        sos = {
            "time":  now,
            "ctemp": safe_float(parts[1]),
            "chum":  safe_float(parts[2]),
            "cpres": safe_float(parts[3]),
            "calt":  safe_float(parts[4]),
            "clat":  safe_float(parts[5]),
            "clon":  safe_float(parts[6]),
            "bpm":   safe_str(parts[7]),
            "slat":  safe_float(parts[8]),
            "slon":  safe_float(parts[9]),
            "stemp": safe_float(parts[10]),
            "shum":  safe_float(parts[11]),
        }

        with lock:
            latest_sos = dict(sos)
            sos_events.append(dict(sos))

        broadcast_sync({"type": "sos_event", **sos})
        return

    # Ignore boot messages, [ERR] lines, "---" separators, etc.

# ── SERIAL READER THREAD ─────────────────────────────────────────────────────
def read_serial():
    try:
        ser = serial.Serial(PORT, BAUD, timeout=1)
        print(f"[OK] Connected to {PORT}")
    except serial.SerialException as e:
        print(f"[ERROR] Could not open {PORT}: {e}")
        print("Available ports:")
        for p in serial.tools.list_ports.comports():
            print(f"  {p.device} — {p.description}")
        print("[INFO] Falling back to DEMO MODE")
        threading.Thread(target=demo_thread, daemon=True).start()
        return

    while True:
        try:
            raw  = ser.readline()
            line = raw.decode("utf-8", errors="replace").strip()
            if line:
                print(f"[SERIAL] {line}")   # echo to console for debugging
                parse_and_dispatch(line, get_timestamp())
        except Exception as e:
            print(f"[SERIAL ERR] {e}")

# ── DEMO MODE ─────────────────────────────────────────────────────────────────
#  Simulates the exact same line formats the ground station outputs
#  so you can test the dashboard without hardware.
def demo_thread():
    import random

    t           = 0
    current_alt = 750.0
    base_lat    = 49.2827
    base_lon    = -123.1207

    while True:
        time.sleep(1.0)
        t += 1

        if current_alt > 0:
            current_alt = max(0.0, current_alt - random.uniform(3.5, 7.0))

        temp     = round(22.0 - current_alt * 0.0065 + random.uniform(-0.3, 0.3), 1)
        pressure = round(1013.25 * math.exp(-current_alt / 8500) + random.uniform(-0.4, 0.4), 1)
        humidity = round(58.0 + math.sin(t * 0.15) * 8 + random.uniform(-0.5, 0.5), 1)
        lat      = round(base_lat + t * 0.00005 + random.uniform(-0.00002, 0.00002), 6)
        lon      = round(base_lon + t * 0.00003 + random.uniform(-0.00002, 0.00002), 6)
        now      = get_timestamp()

        # Normal telemetry — matches ground station CANSAT line exactly
        parse_and_dispatch(
            f"CANSAT,{temp},{humidity},{pressure},{current_alt:.2f},{lat},{lon}",
            now
        )

        # Simulate an SOS_EVENT after t=20
        if t == 20 or (t > 20 and t % 10 == 0):
            bpm   = random.randint(75, 95)
            slat  = round(lat + random.uniform(-0.001, 0.001), 6)
            slon  = round(lon + random.uniform(-0.001, 0.001), 6)
            stemp = round(temp + random.uniform(-1.0, 1.0), 1)
            shum  = round(humidity + random.uniform(-2.0, 2.0), 1)
            parse_and_dispatch(
                f"SOS_EVENT,{temp},{humidity},{pressure},{current_alt:.2f},{lat},{lon},"
                f"{bpm},{slat},{slon},{stemp},{shum}",
                now
            )

# ── EXCEL EXPORT ──────────────────────────────────────────────────────────────
def export_excel():
    with lock:
        log_s = list(data_log)
        sos_s = list(sos_events)

    wb = openpyxl.Workbook()

    # ── Telemetry sheet ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Telemetry"
    headers = ["Time", "Temp (C)", "Humidity (%)", "Pressure (hPa)", "Altitude (m)", "Lat", "Lon"]
    hfill = PatternFill("solid", fgColor="2C3E50")
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = 18

    for ri, e in enumerate(log_s, 2):
        ws.cell(ri, 1, e["time"])
        ws.cell(ri, 2, e["temp"])
        ws.cell(ri, 3, e["humidity"])
        ws.cell(ri, 4, e["pressure"])
        ws.cell(ri, 5, e["altitude"])
        ws.cell(ri, 6, e.get("lat") or "")
        ws.cell(ri, 7, e.get("lon") or "")
        if ri % 2 == 0:
            for col in range(1, 8):
                ws.cell(ri, col).fill = PatternFill("solid", fgColor="ECF0F1")

    # ── SOS Events sheet ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("SOS Events")
    sos_headers = [
        "Time",
        "CanSat Temp (C)", "CanSat Hum (%)", "CanSat Pres (hPa)", "CanSat Alt (m)",
        "CanSat Lat", "CanSat Lon",
        "Hiker BPM", "Hiker Lat", "Hiker Lon", "Hiker Temp (C)", "Hiker Hum (%)"
    ]
    sf = PatternFill("solid", fgColor="922B21")
    for col, h in enumerate(sos_headers, 1):
        c = ws2.cell(1, col, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = sf
        c.alignment = Alignment(horizontal="center")
        ws2.column_dimensions[c.column_letter].width = 20

    for ri, s in enumerate(sos_s, 2):
        ws2.cell(ri,  1, s["time"])
        ws2.cell(ri,  2, s.get("ctemp")  or "")
        ws2.cell(ri,  3, s.get("chum")   or "")
        ws2.cell(ri,  4, s.get("cpres")  or "")
        ws2.cell(ri,  5, s.get("calt")   or "")
        ws2.cell(ri,  6, s.get("clat")   or "")
        ws2.cell(ri,  7, s.get("clon")   or "")
        ws2.cell(ri,  8, s.get("bpm")    or "N/A")
        ws2.cell(ri,  9, s.get("slat")   or "")
        ws2.cell(ri, 10, s.get("slon")   or "")
        ws2.cell(ri, 11, s.get("stemp")  or "")
        ws2.cell(ri, 12, s.get("shum")   or "")
        if ri % 2 == 0:
            for col in range(1, 13):
                ws2.cell(ri, col).fill = PatternFill("solid", fgColor="FADBD8")

    fname = f"cansat_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(fname)
    return os.path.abspath(fname)

# ── WEBSOCKET HANDLER ─────────────────────────────────────────────────────────
async def handler(ws):
    with lock:
        connected_clients.add(ws)
        history = list(data_log)
        sos     = dict(latest_sos) if latest_sos else None

    try:
        # Replay history to newly connected client
        for entry in history:
            await ws.send(json.dumps({"type": "telemetry", **entry}))
        if sos:
            await ws.send(json.dumps({"type": "sos_event", **sos}))

        async for msg in ws:
            data = json.loads(msg)
            if data.get("type") == "export":
                path = export_excel()
                await ws.send(json.dumps({"type": "export_done", "path": path}))

    except websockets.exceptions.ConnectionClosed:
        pass
    finally:
        with lock:
            connected_clients.discard(ws)

# ── MAIN ──────────────────────────────────────────────────────────────────────
async def main():
    global loop
    loop = asyncio.get_running_loop()
    threading.Thread(target=read_serial, daemon=True).start()
    time.sleep(0.5)
    html_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cansat_dashboard.html")
    webbrowser.open(f"file://{html_path}")
    print(f"[OK] WebSocket server on ws://localhost:{WS_PORT}")
    async with websockets.serve(handler, "localhost", WS_PORT):
        await asyncio.Future()

if __name__ == "__main__":
    asyncio.run(main())
